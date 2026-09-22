import io
import json
import logging
from dataclasses import dataclass, field

import fitz  # PyMuPDF
import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from lxml import etree

from app.ingestion.ocr import ocr_image_bytes

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".xlsx", ".xls", ".json", ".xml", ".txt"}


@dataclass
class LoadResult:
    documents: list[Document] = field(default_factory=list)
    ocr_images_processed: int = 0
    warnings: list[str] = field(default_factory=list)


def load_file(file_name: str, content: bytes) -> LoadResult:
    ext = "." + file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in SUPPORTED_EXTENSIONS:
        return LoadResult(warnings=[f"Unsupported file type skipped: {file_name}"])

    try:
        if ext == ".pdf":
            return _load_pdf(file_name, content)
        if ext == ".docx":
            return _load_docx(file_name, content)
        if ext == ".csv":
            return _load_csv(file_name, content)
        if ext in (".xlsx", ".xls"):
            return _load_excel(file_name, content)
        if ext == ".json":
            return _load_json(file_name, content)
        if ext == ".xml":
            return _load_xml(file_name, content)
        if ext == ".txt":
            return _load_txt(file_name, content)
    except Exception as exc:  # noqa: BLE001
        logger.exception("Failed to parse %s", file_name)
        return LoadResult(warnings=[f"Failed to parse {file_name}: {exc}"])

    return LoadResult(warnings=[f"Unsupported file type skipped: {file_name}"])


def _load_pdf(file_name: str, content: bytes) -> LoadResult:
    result = LoadResult()
    pdf = fitz.open(stream=content, filetype="pdf")
    try:
        for page_index, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            ocr_chunks: list[str] = []
            for image_info in page.get_images(full=True):
                xref = image_info[0]
                try:
                    extracted = pdf.extract_image(xref)
                except Exception:  # noqa: BLE001
                    continue
                image_bytes = extracted.get("image")
                if not image_bytes:
                    continue
                ocr_text = ocr_image_bytes(image_bytes)
                result.ocr_images_processed += 1
                if ocr_text:
                    ocr_chunks.append(ocr_text)

            full_text = text
            if ocr_chunks:
                full_text += "\n\n[Embedded image OCR text]\n" + "\n".join(ocr_chunks)

            if full_text.strip():
                result.documents.append(
                    Document(
                        page_content=full_text,
                        metadata={"source": file_name, "page": page_index, "type": "pdf"},
                    )
                )
    finally:
        pdf.close()
    return result


def _load_docx(file_name: str, content: bytes) -> LoadResult:
    result = LoadResult()
    doc = DocxDocument(io.BytesIO(content))

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    table_texts = []
    for t_idx, table in enumerate(doc.tables):
        rows_text = []
        for row in table.rows:
            rows_text.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows_text:
            table_texts.append(f"[Table {t_idx + 1}]\n" + "\n".join(rows_text))
    if table_texts:
        text += "\n\n" + "\n\n".join(table_texts)

    ocr_chunks: list[str] = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                image_bytes = rel.target_part.blob
            except Exception:  # noqa: BLE001
                continue
            ocr_text = ocr_image_bytes(image_bytes)
            result.ocr_images_processed += 1
            if ocr_text:
                ocr_chunks.append(ocr_text)
    if ocr_chunks:
        text += "\n\n[Embedded image OCR text]\n" + "\n".join(ocr_chunks)

    if text.strip():
        result.documents.append(Document(page_content=text, metadata={"source": file_name, "type": "docx"}))
    return result


def _dataframe_to_documents(df: pd.DataFrame, file_name: str, sheet_name: str | None, rows_per_chunk: int = 25) -> list[Document]:
    docs: list[Document] = []
    columns = list(df.columns)
    for start in range(0, len(df), rows_per_chunk):
        chunk_df = df.iloc[start : start + rows_per_chunk]
        lines = [f"Columns: {', '.join(str(c) for c in columns)}"]
        for _, row in chunk_df.iterrows():
            lines.append(", ".join(f"{col}={row[col]}" for col in columns))
        metadata = {
            "source": file_name,
            "type": "tabular",
            "row_start": int(start),
            "row_end": int(start + len(chunk_df) - 1),
        }
        if sheet_name:
            metadata["sheet"] = sheet_name
        docs.append(Document(page_content="\n".join(lines), metadata=metadata))
    return docs


def _load_csv(file_name: str, content: bytes) -> LoadResult:
    df = pd.read_csv(io.BytesIO(content))
    return LoadResult(documents=_dataframe_to_documents(df, file_name, sheet_name=None))


def _load_excel(file_name: str, content: bytes) -> LoadResult:
    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
    docs: list[Document] = []
    for sheet_name, df in sheets.items():
        docs.extend(_dataframe_to_documents(df, file_name, sheet_name=sheet_name))
    return LoadResult(documents=docs)


def _flatten_json(obj, prefix: str = "") -> list[str]:
    lines: list[str] = []
    if isinstance(obj, dict):
        for key, value in obj.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            lines.extend(_flatten_json(value, path))
    elif isinstance(obj, list):
        for idx, value in enumerate(obj):
            path = f"{prefix}[{idx}]"
            lines.extend(_flatten_json(value, path))
    else:
        lines.append(f"{prefix}: {obj}")
    return lines


def _load_json(file_name: str, content: bytes) -> LoadResult:
    data = json.loads(content.decode("utf-8"))
    flattened = "\n".join(_flatten_json(data))
    pretty = json.dumps(data, indent=2, ensure_ascii=False)
    text = f"[Structured JSON key/value view]\n{flattened}\n\n[Raw JSON]\n{pretty}"
    return LoadResult(documents=[Document(page_content=text, metadata={"source": file_name, "type": "json"})])


def _load_xml(file_name: str, content: bytes) -> LoadResult:
    root = etree.fromstring(content)
    text_content = " ".join(t.strip() for t in root.itertext() if t.strip())

    def describe(elem, depth=0) -> list[str]:
        lines = [f"{'  ' * depth}<{elem.tag}> attrs={dict(elem.attrib)}"]
        for child in elem:
            lines.extend(describe(child, depth + 1))
        return lines

    structure = "\n".join(describe(root))
    text = f"[XML structure]\n{structure}\n\n[XML text content]\n{text_content}"
    return LoadResult(documents=[Document(page_content=text, metadata={"source": file_name, "type": "xml"})])


def _load_txt(file_name: str, content: bytes) -> LoadResult:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    return LoadResult(documents=[Document(page_content=text, metadata={"source": file_name, "type": "txt"})])
